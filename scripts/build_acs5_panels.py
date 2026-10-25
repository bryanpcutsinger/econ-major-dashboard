"""Step 8: the 2020-24 5-year panels and bundle v1.

Spec: spec/acs5-panels.yaml (written 2026-10-25 before any estimate; ledger E33-E35). Builds
  - same occupation (panel 3): each major's top 10 jobs; economics vs the major vs everyone in the
    job, ages 25-54 reweighted to the job's age mix (D4); graduate switch
  - local (panel 5): Florida and South Florida; economics vs the major vs all majors; 4 switch settings
  - HC Tables 1-3 rolled forward to 2020-24, in 2024 dollars
and, for the instructor calls, the variants (G1 South Florida PUMAs, L1 local ages, P1 pay measure,
D4 check ages 35-45). Every cell: n, estimate, SDR SE and MOE from PWGTP1-80, hidden flag + reason.

Run from the repo root (after scripts/extract_acs5.py):  python3 scripts/build_acs5_panels.py
Reads cache/acs5-2020-2024-ba.parquet, spec/, raw/benchmarks/ (read-only).
Writes data/bundles/<bundle id>/, data/current.json, output/checks/acs5-variants.json.
"""
import csv
import hashlib
import json
import subprocess
import sys
import time
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_rank_panel import medians  # noqa: E402  (one median rule for the page)
from hc_tables12_check import BLOCKS, recode_block  # noqa: E402

REPS = [f"PWGTP{r}" for r in range(1, 81)]
WCOLS = ["PWGTP"] + REPS                      # column 0 = full-sample weight
Z90 = 1.645
ECON = "Economics"
SF_COUNTIES = {"086", "011", "099"}           # Miami-Dade, Broward, Palm Beach
EXTRA_SF_PUMA = "08625"                       # Monroe County + 10 Miami-Dade tracts (G1 variant)
STATS = {"neg_weight_cells": 0}


# ---------- cells, differences, labels (spec: common) ----------

def sdr(t):
    return float(np.sqrt(4 / 80 * ((t[1:] - t[0]) ** 2).sum()))


def cell(v, W):
    """Median cell: lower median shown, interpolated median for the SE (negative replicate weights = 0)."""
    n = len(v)
    if n == 0:
        return {"n": 0, "hidden": True, "reason": "no people"}, None
    STATS["neg_weight_cells"] += bool((W < 0).any())
    lower, interp = medians(v, np.maximum(W, 0))
    est, se = float(lower[0]), sdr(interp)
    reason = "n < 100" if n < 100 else ("MOE / estimate > 0.30" if Z90 * se / est > 0.30 else "")
    if not np.isfinite(se):
        reason = "replicate weights sum to 0"
    out = {"n": n, "median": round(est, 2), "se": round(se, 2), "moe": round(Z90 * se, 2),
           "hidden": bool(reason), "reason": reason}
    return out, interp


def compare(a, b):
    """Economics (a) minus other (b), inside each replicate; label only if both cells show."""
    (ca, ta), (cb, tb) = a, b
    if ca["hidden"] or cb["hidden"]:
        return {"label": None}
    d = ta - tb
    se = sdr(d)
    lab = "better" if d[0] > Z90 * se else ("worse" if d[0] < -Z90 * se else "about the same")
    return {"d": round(float(d[0]), 2), "se": round(se, 2), "moe": round(Z90 * se, 2), "label": lab}


def ratio(num, den):
    """Weighted share (81 columns) with its SDR SE."""
    r = num / den
    return {"share": round(float(r[0]), 6), "se": round(sdr(r), 6), "moe": round(Z90 * sdr(r), 6)}


# ---------- age reweighting (spec: same_occupation.age_reweighting, D4) ----------

def reweight(W, age, target_age, target_w):
    """Weight a group to the target's age mix. Factors from PWGTP only, applied to all 81 columns.

    Bins of 5 years (25-29 ... 50-54), then of 10 (25-34, 35-44, 45-54); None if a target bin is empty
    in the group at both widths (cell hidden: 'age mix').
    """
    for width in (5, 10):
        nb = 30 // width
        tb = np.bincount((target_age - 25) // width, weights=target_w, minlength=nb)
        gb = np.bincount((age - 25) // width, weights=W[:, 0], minlength=nb)
        if np.any((tb > 0) & (gb <= 0)):
            continue
        f = np.where(gb > 0, (tb / tb.sum()) / np.where(gb > 0, gb / gb.sum(), 1), 0)
        return W * f[(age - 25) // width][:, None]
    return None


def age_cell(d, idx, tgt_idx, reweight_on):
    v, W, age = d["pay"][idx], d["W"][idx], d["AGEP"][idx]
    if reweight_on and len(idx):
        W2 = reweight(W, age, d["AGEP"][tgt_idx], d["W"][tgt_idx, 0])
        if W2 is None:
            return {"n": len(idx), "hidden": True, "reason": "age mix"}, None
        W = W2
    return cell(v, W)


# ---------- data ----------

def load():
    df = pd.read_parquet(ROOT / "cache" / "acs5-2020-2024-ba.parquet")
    m73 = {int(r["fod1p"]): r["nyfed_major"] for r in csv.DictReader(open(ROOT / "spec" / "major73-hc.csv"))
           if r["nyfed_major"]}
    nyfed = [r["Major"] for r in csv.DictReader(
        open(ROOT / "raw/benchmarks/nyfed-2024/college-labor-outcomes-by-major-data.csv", encoding="utf-8-sig"))]
    majors = sorted(m for m in nyfed if m != "Overall")
    df["g"] = df.FOD1P.map(lambda c: m73.get(int(c)))
    adj = df.ADJINC.to_numpy(float) / 1e6
    df["pay_wagp"] = df.WAGP.to_numpy(float) * adj        # D2: 2024 dollars
    df["pay_pernp"] = df.PERNP.to_numpy(float) * adj
    return df, majors


def arrays(s, pay_col):
    """Numpy views of one sample, sorted as the extract (SERIALNO, SPORDER)."""
    return {"pay": s[pay_col].to_numpy(float), "W": s[WCOLS].to_numpy(float), "AGEP": s.AGEP.to_numpy(int),
            "g": s.g.to_numpy(object), "OCCP": s.OCCP.astype(int).to_numpy()}


def base_mask(df, grad, fulltime, ages, pay_col, threshold=5500, strict=False):
    m = (df.SCH == 1) & df.ESR.isin([1, 2, 4, 5]) & df.AGEP.between(*ages)
    m &= (df.SCHL >= 21) if grad else (df.SCHL == 21)
    if fulltime:
        m &= (df.WKHP >= 35) & (df.WKWN >= 40)
    m &= (df[pay_col] > threshold) if strict else (df[pay_col] >= threshold)
    return m.to_numpy()


# ---------- panel 3: same occupation ----------

def index_by(keys):
    """{key: row indices} for an array of keys."""
    o = np.argsort(keys, kind="mergesort")
    u, starts = np.unique(keys[o], return_index=True)
    return dict(zip(u.tolist(), np.split(o, starts[1:])))


def top_jobs(d, idx):
    """Top 10 OCCP by weighted count (ties -> lower code)."""
    w = pd.Series(d["W"][idx, 0]).groupby(d["OCCP"][idx]).sum()
    return sorted(w.index.tolist(), key=lambda o: (-w[o], o))[:10]


def same_occupation(df, majors, occ_label, grad, ages, reweight_on, pay_col, rows_from=None):
    s = df[base_mask(df, grad, True, ages, pay_col)].reset_index(drop=True)
    d = arrays(s, pay_col)
    by_g = {m: np.flatnonzero(d["g"] == m) for m in majors}
    by_o = index_by(d["OCCP"])
    none = np.array([], int)
    tops = rows_from or {m: top_jobs(d, by_g[m]) for m in majors}
    cache = {}

    def get(grp, o):
        key = (grp, o)
        if key not in cache:
            tgt = by_o.get(o, none)
            if grp == "all":
                cache[key] = cell(d["pay"][tgt], d["W"][tgt])
            else:
                idx = tgt[d["g"][tgt] == grp]
                cache[key] = age_cell(d, idx, tgt, reweight_on)
        return cache[key]

    out = {}
    for m in majors:
        tot = d["W"][by_g[m]].sum(axis=0)
        lists = {}
        for which, rows in (("major_top", tops[m]), ("economics_top", tops[ECON])):
            lst = []
            for o in rows:
                e, a = get(ECON, o), get("all", o)
                row = {"occp": int(o), "label": occ_label[int(o)],
                       "major_share": ratio(d["W"][by_o.get(o, none)[d["g"][by_o.get(o, none)] == m]].sum(axis=0), tot),
                       "economics": e[0], "all": a[0], "econ_vs_all": compare(e, a)}
                if m != ECON:
                    mm = get(m, o)
                    row["major"] = mm[0]
                    row["econ_vs_major"] = compare(e, mm)
                lst.append(row)
            key = "econ_vs_major" if m != ECON else "econ_vs_all"
            labs = [r[key]["label"] for r in lst]
            lists[which] = {"rows": lst, "headline": {k: labs.count(k) for k in ("better", "about the same", "worse")},
                            "hidden_rows": labs.count(None)}
        out[m] = {"n": int(len(by_g[m])), **lists}
    present = sum(len(by_g[m]) > 0 for m in majors)
    return out, tops, present


# ---------- panel 5: local ----------

def local(df, majors, areas, grad, fulltime, ages, reweight_on, pay_col):
    out = {}
    for area, amask in areas.items():
        s = df[base_mask(df, grad, fulltime, ages, pay_col) & amask].reset_index(drop=True)
        d = arrays(s, pay_col)
        everyone = np.arange(len(s))
        a = cell(d["pay"], d["W"])
        e = age_cell(d, np.flatnonzero(d["g"] == ECON), everyone, reweight_on)
        rows = {}
        for m in majors:
            mm = age_cell(d, np.flatnonzero(d["g"] == m), everyone, reweight_on) if m != ECON else e
            rows[m] = {"major": mm[0], "econ_vs_major": compare(e, mm) if m != ECON else {"label": None}}
        out[area] = {"economics": e[0], "all": a[0], "econ_vs_all": compare(e, a), "by_major": rows,
                     "majors_present": sum(r["major"]["n"] > 0 for r in rows.values())}
    return out


# ---------- HC Tables 1-3 rolled forward ----------

def industry_map():
    """HC industry_category on INDP, plus the 33 new codes via the Census 2017 -> 2022 changes (E34)."""
    hc = recode_block(*BLOCKS["table2"][:2])
    wb = openpyxl.load_workbook(ROOT / "raw/benchmarks/census-codelists/2022-Census-Industry-Code-List-with-Crosswalk.xlsx",
                                read_only=True)
    pred = {}
    for row in list(wb["2022 Ind Code Changes"].iter_rows(values_only=True))[3:]:
        if isinstance(row[0], int) and isinstance(row[2], int):
            pred.setdefault(row[2], set()).add(row[0])
    for row in wb["Explanation of 2022 Changes "].iter_rows(values_only=True):
        old, new = row[1], row[5]
        if str(old).strip().isdigit() and str(new).strip().isdigit() and int(old) != int(new):
            pred.setdefault(int(new), set()).add(int(old))
    ext = dict(hc)
    for new, olds in pred.items():
        groups = {hc[o] for o in olds if o in hc}
        if new not in hc and len(groups) == 1:
            ext[new] = groups.pop()
    return ext


def hc_tables(df, majors, occ_label, ind_map):
    m = base_mask(df, False, True, (35, 45), "pay_pernp", threshold=5499, strict=True)
    s = df[m].reset_index(drop=True)
    d = arrays(s, "pay_pernp")
    out = {}
    econ = np.flatnonzero(d["g"] == ECON)
    econ_tot = d["W"][econ].sum(axis=0)
    indp = s.INDP.astype(int).to_numpy()
    unmapped_ind = sorted({int(c) for c in indp if c not in ind_map})
    for key, codes, labels, top in (("table1", d["OCCP"], occ_label, 40), ("table2", indp, ind_map, 20)):
        grp = np.array([labels.get(int(c), f"code {c}") for c in codes], dtype=object)
        freq = pd.Series(grp[econ]).value_counts()
        order = sorted(freq.index, key=lambda g: (-freq[g], g))[:top]
        rows = []
        for g in order:
            idx = econ[grp[econ] == g]
            c = cell(d["pay"][idx], d["W"][idx])[0]
            rows.append({"group": g, "frequency": int(freq[g]), "share_unweighted": round(100 * freq[g] / len(econ), 4),
                         "share_weighted": ratio(d["W"][idx].sum(axis=0), econ_tot), "median": c})
        out[key] = {"economics_n": int(len(econ)), "rows": rows}
    total = len(s)
    t3 = []
    selfemp = s.COW.isin([6, 7]).to_numpy()
    all_tot = d["W"].sum(axis=0)
    for mj in majors:
        idx = np.flatnonzero(d["g"] == mj)
        k = int(selfemp[idx].sum())
        se_rate = ratio(d["W"][idx[selfemp[idx]]].sum(axis=0), d["W"][idx].sum(axis=0)) if len(idx) else None
        hide_rate = len(idx) < 100 or min(k, len(idx) - k) < 10
        t3.append({"major": mj, "frequency": int(len(idx)), "share_unweighted": round(100 * len(idx) / total, 4),
                   "share_weighted": ratio(d["W"][idx].sum(axis=0), all_tot),
                   "median": cell(d["pay"][idx], d["W"][idx])[0],
                   "self_employment": {"unweighted": round(100 * k / len(idx), 4) if len(idx) else None,
                                       "weighted": se_rate, "hidden": hide_rate}})
    out["table3"] = {"total_n": total, "rows": t3}
    return out, unmapped_ind


# ---------- bundle ----------

def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for b in iter(lambda: fh.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def dump(obj):
    return json.dumps(obj, indent=1, sort_keys=True, ensure_ascii=False) + "\n"


def main():
    t0 = time.time()
    df, majors = load()
    occ_label = recode_block(*BLOCKS["table1"][:2])
    ind_map = industry_map()

    # Gates (spec: gates)
    fod_unmapped = df.g.isna() & (df.FOD1P != 0)
    gates = {"fod1p_unmapped_persons": int(fod_unmapped.sum()),
             "occp_without_label": sorted({int(c) for c in df.OCCP.dropna()} - set(occ_label)),
             "indp_unmapped": sorted({int(c) for c in df.INDP.dropna()} - set(ind_map))}
    tract = pd.read_csv(ROOT / "raw/benchmarks/census-geo/2020_Census_Tract_to_2020_PUMA.txt", dtype=str,
                        encoding="utf-8-sig")
    fl = tract[tract.STATEFP == "12"]
    inside = sorted(p for p, cs in fl.groupby("PUMA5CE").COUNTYFP.apply(set).items() if cs <= SF_COUNTIES)
    gates["south_florida_pumas"] = len(inside)
    print("gates:", gates, flush=True)

    in_fl = (df.STATE == "12").to_numpy()
    areas = {"florida": in_fl, "south_florida": in_fl & df.PUMA.isin(inside).to_numpy()}

    bundle, variants = {}, {}
    # Panel 3, default (D4: ages 25-54 reweighted; wages)
    so, tops, present = {}, {}, {}
    for grad in (False, True):
        k = f"grad={int(grad)}"
        so[k], tops[k], present[k] = same_occupation(df, majors, occ_label, grad, (25, 54), True, "pay_wagp")
        print(f"  same occupation {k}: {time.time() - t0:.0f} s", flush=True)
    gates["same_occupation_majors_present"] = present
    bundle["same_occupation"] = so
    # D4 check (35-45, no reweighting) and P1 (earnings), on the default rows, graduate switch off
    variants["same_occupation_ages_35_45"] = same_occupation(df, majors, occ_label, False, (35, 45), False,
                                                             "pay_wagp", rows_from=tops["grad=0"])[0]
    variants["same_occupation_pernp"] = same_occupation(df, majors, occ_label, False, (25, 54), True,
                                                        "pay_pernp", rows_from=tops["grad=0"])[0]
    print(f"  same occupation variants: {time.time() - t0:.0f} s", flush=True)

    # Panel 5, default (L1: 25-54 reweighted; G1: 52 PUMAs; P1: wages), 4 switch settings
    loc = {}
    for grad in (False, True):
        for ft in (True, False):
            loc[f"grad={int(grad)}|fulltime={int(ft)}"] = local(df, majors, areas, grad, ft, (25, 54), True, "pay_wagp")
    bundle["local"] = loc
    dflt = "grad=0|fulltime=1"
    variants["local_ages_35_45"] = local(df, majors, areas, False, True, (35, 45), False, "pay_wagp")
    variants["local_pernp"] = local(df, majors, areas, False, True, (25, 54), True, "pay_pernp")
    variants["local_sf_with_08625"] = local(df, majors, {"south_florida": in_fl & df.PUMA.isin(inside + [EXTRA_SF_PUMA]).to_numpy()},
                                            False, True, (25, 54), True, "pay_wagp")
    print(f"  local: {time.time() - t0:.0f} s", flush=True)

    # HC Tables 1-3, 2020-24
    bundle["hc_tables_5yr"], gates["hc_table2_indp_unmapped"] = hc_tables(df, majors, occ_label, ind_map)
    gates["hc_table3_majors_present"] = sum(r["frequency"] > 0 for r in bundle["hc_tables_5yr"]["table3"]["rows"])

    # Cross-check (warning only): national economics mid-career pay, rank-panel rules, 5-year vs 1-year
    s = df[base_mask(df, False, True, (35, 45), "pay_wagp") & (df.g == ECON).to_numpy()]
    c5 = cell(s.pay_wagp.to_numpy(float), s[WCOLS].to_numpy(float))[0]
    rank = json.loads((ROOT / "output/panels/rank-2024.json").read_text())
    r1 = next(r for r in rank["measures"]["pay_mid|grad=0|fulltime=1"] if r["major"] == ECON)
    variants["cross_check_econ_mid_pay"] = {"acs5_2020_24": c5, "acs1_2024_rank_panel": {"estimate": r1["estimate"], "se": r1["se"], "n": r1["n"]}}
    gates["negative_weight_cells"] = STATS["neg_weight_cells"]

    # Bundle v1 (spec: bundle_v1)
    codes = {"FOD1P": sorted({int(c) for c in df.FOD1P}), "OCCP": sorted({int(c) for c in df.OCCP.dropna()}),
             "INDP": sorted({int(c) for c in df.INDP.dropna()}),
             "PUMA": sorted(set(df.STATE + df.PUMA))}
    codehash = hashlib.sha256(json.dumps(codes, sort_keys=True).encode()).hexdigest()[:8]
    bid = f"2024-1yr_2020-24-5yr_{codehash}"
    out = ROOT / "data" / "bundles" / bid
    out.mkdir(parents=True, exist_ok=True)
    (out / "rank.json").write_bytes((ROOT / "output/panels/rank-2024.json").read_bytes())
    for name in ("same_occupation", "local", "hc_tables_5yr"):
        (out / f"{name}.json").write_text(dump(bundle[name]))
    git = subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    dirty = bool(subprocess.run(["git", "status", "--porcelain", "scripts", "spec"], cwd=ROOT,
                                capture_output=True, text=True).stdout.strip())
    manifest = {"bundle": bid, "git_commit": git, "worktree_dirty": dirty, "gates": gates,
                "dollars": "2024 dollars (ADJINC)", "years": {"rank": "ACS 2024 1-year", "other": "ACS 2020-24 5-year"},
                "sources": {p: sha(ROOT / p) for p in ["raw/acs1-2024/csv_pus.zip", "raw/acs5-2020-2024/csv_pus.zip"]},
                "spec": {p: sha(ROOT / p) for p in ["spec/acs5-panels.yaml", "spec/rank-panel.yaml", "spec/major73-hc.csv"]},
                "code": {p: sha(ROOT / p) for p in ["scripts/build_acs5_panels.py", "scripts/build_rank_panel.py",
                                                     "scripts/hc_tables12_check.py", "scripts/extract_acs5.py"]},
                "files": {f.name: sha(f) for f in sorted(out.glob("*.json")) if f.name != "manifest.json"}}
    (out / "manifest.json").write_text(dump(manifest))
    (ROOT / "data" / "current.json").write_text(dump({"bundle": bid}))
    (ROOT / "output" / "checks" / "acs5-variants.json").write_text(dump(variants))
    print(f"wrote data/bundles/{bid}/ ({len(manifest['files']) + 1} files); {time.time() - t0:.0f} s")
    print("gates:", json.dumps(gates))


if __name__ == "__main__":
    main()
